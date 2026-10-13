"""
Build 11-slide presentation for Task Management System.
Requires: pip install python-pptx
Run: python docs/presentation/build_task_management_deck.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "TaskManagementSystem_Presentation.pptx"
ASSETS = ROOT / "assets"

NAVY = RGBColor(0x0F, 0x2A, 0x44)
NAVY_MID = RGBColor(0x1E, 0x3A, 0x5F)
TEAL = RGBColor(0x0D, 0x94, 0x88)
TEAL_LIGHT = RGBColor(0xCC, 0xFB, 0xF1)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF8, 0xFA, 0xFC)
CARD = RGBColor(0xFF, 0xFF, 0xFF)

TEAM_MEMBERS = [
    ("Denish Murawala", "Team Lead / Backend", ASSETS / "denish.jpg"),
    ("Avinash Kanaujia", "Frontend / UI", ASSETS / "avinash.jpg"),
    ("Anushtup Ganguly", "Services / Integration", ASSETS / "anushtup.jpg"),
]


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_accent_bar(slide, top=0, height=0.12, color=TEAL) -> None:
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(top), Inches(13.333), Inches(height)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def add_footer(slide, text: str = "Task Management System | .NET Microservices") -> None:
    foot = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(7.15), Inches(13.333), Inches(0.35)
    )
    foot.fill.solid()
    foot.fill.fore_color.rgb = NAVY
    foot.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(7.18), Inches(12), Inches(0.3))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0xCB, 0xE8, 0xE6)


def add_section_header(slide, title: str, subtitle: str = "") -> None:
    set_slide_bg(slide, LIGHT)
    add_accent_bar(slide, 0, 0.1, TEAL)
    header = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0.1), Inches(13.333), Inches(1.15)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12), Inches(0.65))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        sb = slide.shapes.add_textbox(Inches(0.55), Inches(1.35), Inches(12), Inches(0.45))
        sp = sb.text_frame.paragraphs[0]
        sp.text = subtitle
        sp.font.size = Pt(15)
        sp.font.color.rgb = MUTED
    add_footer(slide)


def add_team_photos_row(
    slide,
    start_left: float,
    photo_top: float,
    name_color: RGBColor = WHITE,
    pic_size_in: float = 1.35,
    column_spacing: float = 2.15,
) -> None:
    """Same layout on title and thank-you slides: photo strip + full names below."""
    pic_size = Inches(pic_size_in)
    name_w = pic_size_in + 0.55
    for i, (name, _role, photo) in enumerate(TEAM_MEMBERS):
        left = start_left + i * column_spacing
        if photo.is_file():
            slide.shapes.add_picture(str(photo), Inches(left), Inches(photo_top), pic_size, pic_size)
        else:
            ph = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.OVAL, Inches(left), Inches(photo_top), pic_size, pic_size
            )
            ph.fill.solid()
            ph.fill.fore_color.rgb = TEAL_LIGHT
            ph.line.color.rgb = TEAL

        name_top = photo_top + pic_size_in + 0.12
        nm = slide.shapes.add_textbox(
            Inches(left - 0.2), Inches(name_top), Inches(name_w), Inches(0.75)
        )
        tf = nm.text_frame
        tf.word_wrap = True
        nmp = tf.paragraphs[0]
        nmp.text = name
        nmp.font.size = Pt(11)
        nmp.font.bold = True
        nmp.font.color.rgb = name_color
        nmp.alignment = PP_ALIGN.CENTER


def add_title_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)
    add_accent_bar(slide, 0, 0.14, TEAL)

    # decorative panel
    panel = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0.14), Inches(7.2), Inches(7.5)
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = NAVY_MID
    panel.line.fill.background()

    title = slide.shapes.add_textbox(Inches(0.65), Inches(1.0), Inches(6.2), Inches(1.4))
    tp = title.text_frame.paragraphs[0]
    tp.text = "Task Management System"
    tp.font.size = Pt(40)
    tp.font.bold = True
    tp.font.color.rgb = WHITE

    sub = slide.shapes.add_textbox(Inches(0.65), Inches(2.35), Inches(6.2), Inches(1.0))
    sp = sub.text_frame
    sp.word_wrap = True
    p0 = sp.paragraphs[0]
    p0.text = "Microservices Architecture with .NET 10 & React"
    p0.font.size = Pt(18)
    p0.font.color.rgb = TEAL_LIGHT
    p1 = sp.add_paragraph()
    p1.text = "Academic Project Presentation"
    p1.font.size = Pt(14)
    p1.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    p1.space_before = Pt(6)

    chips = slide.shapes.add_textbox(Inches(0.65), Inches(3.5), Inches(6.0), Inches(0.8))
    cp = chips.text_frame.paragraphs[0]
    cp.text = "API Gateway  •  JWT Auth  •  Kanban  •  Docker"
    cp.font.size = Pt(12)
    cp.font.color.rgb = RGBColor(0xCB, 0xE8, 0xE6)

    team_lbl = slide.shapes.add_textbox(Inches(7.35), Inches(0.55), Inches(5.5), Inches(0.4))
    tlp = team_lbl.text_frame.paragraphs[0]
    tlp.text = "PROJECT TEAM"
    tlp.font.size = Pt(12)
    tlp.font.bold = True
    tlp.font.color.rgb = TEAL_LIGHT

    add_team_photos_row(slide, start_left=7.35, photo_top=1.05, name_color=WHITE)

    add_footer(slide)


def slide_overview(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Project Overview", "What we built and why it matters")

    highlights = [
        ("Collaboration", "Projects, sprints, tasks, comments & files in one place"),
        ("Microservices", "8 focused .NET services behind an Ocelot gateway"),
        ("Roles", "Admin configures; users execute work on Kanban boards"),
    ]
    for i, (title, desc) in enumerate(highlights):
        left = 0.55 + i * 4.15
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(1.95),
            Inches(3.85),
            Inches(2.0),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = TEAL
        card.line.width = Pt(2)
        accent = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(1.95), Inches(3.85), Inches(0.12)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = TEAL
        accent.line.fill.background()
        tb = slide.shapes.add_textbox(Inches(left + 0.2), Inches(2.2), Inches(3.45), Inches(1.6))
        tf = tb.text_frame
        tf.word_wrap = True
        h = tf.paragraphs[0]
        h.text = title
        h.font.bold = True
        h.font.size = Pt(17)
        h.font.color.rgb = NAVY
        b = tf.add_paragraph()
        b.text = desc
        b.font.size = Pt(12)
        b.font.color.rgb = SLATE
        b.space_before = Pt(8)

    note = slide.shapes.add_textbox(Inches(0.55), Inches(4.25), Inches(12.2), Inches(2.5))
    nf = note.text_frame
    nf.word_wrap = True
    lines = [
        "End-to-end task board for software teams — similar spirit to Jira/Trello, scoped for learning.",
        "Runs fully on your machine: Docker (SQL Server, MongoDB, Redis) + React UI on port 5173.",
        "Demo dataset: Acme Digital — 3 products, active sprints, Kanban flow Todo → Done.",
    ]
    for i, line in enumerate(lines):
        p = nf.paragraphs[0] if i == 0 else nf.add_paragraph()
        p.text = f"▸  {line}"
        p.font.size = Pt(14)
        p.font.color.rgb = SLATE
        p.space_after = Pt(6)


def slide_problem_solution(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Problem & Solution", "From scattered work to a governed platform")

    blocks = [
        ("THE PROBLEM", NAVY, RGBColor(0xFE, 0xE2, 0xE2), [
            "Tasks live in chat & spreadsheets — no single source of truth",
            "Sprints and ownership are unclear",
            "Everyone can change structure — or no one can",
        ]),
        ("OUR SOLUTION", TEAL, TEAL_LIGHT, [
            "One API gateway + JWT-secured microservices",
            "Projects, sprints, Kanban, notifications",
            "Admin vs User roles; assignee gets project access",
        ]),
    ]
    for idx, (heading, hdr_color, bullet_bg, points) in enumerate(blocks):
        left = 0.5 + idx * 6.35
        w = Inches(6.0)
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.85), w, Inches(4.85)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        hdr = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(1.85), w, Inches(0.65)
        )
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = hdr_color
        hdr.line.fill.background()
        ht = slide.shapes.add_textbox(Inches(left + 0.25), Inches(1.95), w - Inches(0.5), Inches(0.5))
        hp = ht.text_frame.paragraphs[0]
        hp.text = heading
        hp.font.size = Pt(16)
        hp.font.bold = True
        hp.font.color.rgb = WHITE
        y = 2.65
        for pt in points:
            row = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                Inches(left + 0.25),
                Inches(y),
                w - Inches(0.5),
                Inches(0.85),
            )
            row.fill.solid()
            row.fill.fore_color.rgb = bullet_bg if idx == 1 else RGBColor(0xF1, 0xF5, 0xF9)
            row.line.fill.background()
            tb = slide.shapes.add_textbox(
                Inches(left + 0.4), Inches(y + 0.12), w - Inches(0.7), Inches(0.7)
            )
            tb.text_frame.word_wrap = True
            bp = tb.text_frame.paragraphs[0]
            bp.text = pt
            bp.font.size = Pt(13)
            bp.font.color.rgb = SLATE
            y += 1.0


def slide_architecture(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "High-Level Architecture", "Client → Gateway → Microservices → Databases")

    def box(left, top, w, h, text, fill=WHITE, text_color=NAVY):
        sh = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h)
        )
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        sh.line.color.rgb = NAVY_MID
        sh.line.width = Pt(1)
        tf = sh.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = text_color if fill == TEAL else NAVY
        p.alignment = PP_ALIGN.CENTER
        return sh

    ui = box(5.35, 1.9, 2.6, 0.6, "React SPA\nlocalhost:5173", RGBColor(0xE0, 0xF2, 0xFE))
    gw = box(5.05, 2.75, 3.2, 0.65, "Ocelot API Gateway\nlocalhost:5204", TEAL, WHITE)
    slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(6.65), Inches(2.5), Inches(6.65), Inches(2.75)
    ).line.color.rgb = MUTED
    services = [
        "Auth", "Project", "Task", "Comment",
        "File", "Notification", "Dashboard", "AI",
    ]
    for i, name in enumerate(services):
        col, row = i % 4, i // 4
        box(0.45 + col * 3.15, 3.65 + row * 0.9, 2.85, 0.72, name)
        slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(6.65),
            Inches(3.4),
            Inches(1.85 + col * 3.15 + 1.4),
            Inches(3.65 + row * 0.9),
        ).line.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
    box(3.8, 5.55, 5.7, 0.55, "SQL Server  |  MongoDB (comments)  |  Redis (SignalR)")


def slide_tech_stack(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Technology Stack")
    rows = [
        ("Frontend", "React, Vite, Axios, SignalR client"),
        ("API", "ASP.NET Core 10, Ocelot Gateway, JWT auth"),
        ("Services", "Auth, Project, Task, Comment, File, Notification, Dashboard, AI"),
        ("Data", "EF Core + SQL Server; MongoDB (comments); Redis (SignalR)"),
        ("DevOps", "Docker Compose, GitHub Actions CI, xUnit tests"),
    ]
    tbl = slide.shapes.add_table(len(rows) + 1, 2, Inches(0.6), Inches(1.9), Inches(12.0), Inches(4.2)).table
    hdr = tbl.rows[0].cells
    hdr[0].text = "Layer"
    hdr[1].text = "Technologies"
    for c in hdr:
        for p in c.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(14)
            p.font.color.rgb = WHITE
        c.fill.solid()
        c.fill.fore_color.rgb = NAVY
    for r, (layer, tech) in enumerate(rows, start=1):
        tbl.rows[r].cells[0].text = layer
        tbl.rows[r].cells[1].text = tech
        bg = LIGHT if r % 2 == 0 else WHITE
        for ci in (0, 1):
            tbl.rows[r].cells[ci].fill.solid()
            tbl.rows[r].cells[ci].fill.fore_color.rgb = bg
            for p in tbl.rows[r].cells[ci].text_frame.paragraphs:
                p.font.size = Pt(13)
                p.font.color.rgb = SLATE


def slide_services(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Microservices", "Each service owns one business capability")

    items = [
        ("Auth", "Login, JWT, profiles"),
        ("Project", "Teams, projects, sprints, access"),
        ("Task", "Tasks, subtasks, Kanban, assign"),
        ("Comment", "MongoDB + SignalR hub"),
        ("File", "Per-task uploads"),
        ("Notify", "In-app alerts"),
        ("Dashboard", "Aggregated metrics"),
        ("AI", "Optional Gemini chat"),
    ]
    for i, (name, desc) in enumerate(items):
        col, row = i % 4, i // 4
        left = 0.45 + col * 3.15
        top = 1.95 + row * 2.35
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(2.95),
            Inches(2.05),
        )
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        stripe = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(left), Inches(top), Inches(0.1), Inches(2.05)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = TEAL if i % 2 == 0 else NAVY_MID
        stripe.line.fill.background()
        nt = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.25), Inches(2.5), Inches(0.5))
        np = nt.text_frame.paragraphs[0]
        np.text = name
        np.font.bold = True
        np.font.size = Pt(15)
        np.font.color.rgb = NAVY
        dt = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.85), Inches(2.55), Inches(1.0))
        dt.text_frame.word_wrap = True
        dp = dt.text_frame.paragraphs[0]
        dp.text = desc
        dp.font.size = Pt(11)
        dp.font.color.rgb = MUTED


def slide_request_flow(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Request Flow", "Example: user updates task status on Kanban")

    steps = [
        ("Browser", "Sends PATCH + JWT"),
        ("Gateway", "Routes to Task service"),
        ("Task API", "Checks assignee / creator"),
        ("Project API", "Validates project access"),
        ("UI", "Board refreshes"),
    ]
    y = 2.0
    for i, (title, detail) in enumerate(steps):
        num = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(0.55), Inches(y), Inches(0.45), Inches(0.45))
        num.fill.solid()
        num.fill.fore_color.rgb = TEAL
        num.line.fill.background()
        nt = slide.shapes.add_textbox(Inches(0.55), Inches(y + 0.05), Inches(0.45), Inches(0.4))
        ntp = nt.text_frame.paragraphs[0]
        ntp.text = str(i + 1)
        ntp.font.bold = True
        ntp.font.size = Pt(14)
        ntp.font.color.rgb = WHITE
        ntp.alignment = PP_ALIGN.CENTER
        card = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.15), Inches(y - 0.05), Inches(11.5), Inches(0.75)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        tt = slide.shapes.add_textbox(Inches(1.35), Inches(y + 0.02), Inches(3.5), Inches(0.35))
        ttp = tt.text_frame.paragraphs[0]
        ttp.text = title
        ttp.font.bold = True
        ttp.font.size = Pt(14)
        ttp.font.color.rgb = NAVY
        dt = slide.shapes.add_textbox(Inches(4.9), Inches(y + 0.05), Inches(7.5), Inches(0.4))
        dtp = dt.text_frame.paragraphs[0]
        dtp.text = detail
        dtp.font.size = Pt(12)
        dtp.font.color.rgb = MUTED
        if i < len(steps) - 1:
            slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, Inches(0.77), Inches(y + 0.45), Inches(0.77), Inches(y + 0.85)
            ).line.color.rgb = TEAL
        y += 0.95


def slide_rbac(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Roles & Access Control")
    add_bullets = slide.shapes.add_textbox(Inches(0.55), Inches(1.85), Inches(5.8), Inches(4.5))
    tf = add_bullets.text_frame
    for i, line in enumerate([
        "Admin: teams, projects, sprints, members, assign & delete tasks.",
        "User: Kanban, comments, files on permitted projects.",
        "Access via team membership OR task assignment.",
        "UI hides admin controls for non-admin JWT role.",
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"▸  {line}"
        p.font.size = Pt(14)
        p.font.color.rgb = SLATE
        p.space_after = Pt(8)
    for role, perms, left, color in [
        ("Admin", "Structure & assignment", 6.6, NAVY),
        ("User", "Execution & collaboration", 9.75, TEAL),
    ]:
        c = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(2.2), Inches(3.0), Inches(3.8)
        )
        c.fill.solid()
        c.fill.fore_color.rgb = color
        c.line.fill.background()
        tf2 = c.text_frame
        tf2.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = tf2.paragraphs[0]
        p0.text = role
        p0.font.size = Pt(26)
        p0.font.bold = True
        p0.font.color.rgb = WHITE
        p0.alignment = PP_ALIGN.CENTER
        p1 = tf2.add_paragraph()
        p1.text = perms
        p1.font.size = Pt(12)
        p1.font.color.rgb = TEAL_LIGHT
        p1.alignment = PP_ALIGN.CENTER


def slide_demo_flow(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Live Demo Script", "Follow this order during the class presentation")

    steps = [
        ("Admin login", "admin@local.test — open Projects: Customer Mobile App, Web Portal, Analytics"),
        ("Kanban", "Move task: Todo → InProgress → Testing → Done"),
        ("Assign", "Admin assigns task; assignee gains project access automatically"),
        ("User login", "user@local.test — task detail: subtasks, comments, files"),
        ("Wrap-up", "Dashboard, notifications, mention scripts/seed-presentation-demo.ps1"),
    ]
    for i, (title, body) in enumerate(steps):
        top = 1.9 + i * 1.05
        badge = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(top), Inches(12.3), Inches(0.92)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = WHITE if i % 2 == 0 else RGBColor(0xF1, 0xF5, 0xF9)
        badge.line.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
        step_lbl = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(top + 0.15), Inches(1.1), Inches(0.55)
        )
        step_lbl.fill.solid()
        step_lbl.fill.fore_color.rgb = TEAL
        step_lbl.line.fill.background()
        st = slide.shapes.add_textbox(Inches(0.65), Inches(top + 0.22), Inches(1.1), Inches(0.4))
        stp = st.text_frame.paragraphs[0]
        stp.text = f"Step {i + 1}"
        stp.font.size = Pt(11)
        stp.font.bold = True
        stp.font.color.rgb = WHITE
        stp.alignment = PP_ALIGN.CENTER
        tt = slide.shapes.add_textbox(Inches(1.95), Inches(top + 0.12), Inches(10.5), Inches(0.35))
        ttp = tt.text_frame.paragraphs[0]
        ttp.text = title
        ttp.font.bold = True
        ttp.font.size = Pt(15)
        ttp.font.color.rgb = NAVY
        bt = slide.shapes.add_textbox(Inches(1.95), Inches(top + 0.45), Inches(10.5), Inches(0.4))
        btp = bt.text_frame.paragraphs[0]
        btp.text = body
        btp.font.size = Pt(11)
        btp.font.color.rgb = MUTED


def slide_quality(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_section_header(slide, "Engineering Quality", "How we keep the system reliable")

    metrics = [("130+", "Automated tests"), ("8", "Microservices"), ("CI", "GitHub Actions"), ("1", "Monorepo slnx")]
    for i, (val, lbl) in enumerate(metrics):
        left = 0.55 + i * 3.15
        m = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(1.95), Inches(2.9), Inches(1.35)
        )
        m.fill.solid()
        m.fill.fore_color.rgb = NAVY
        m.line.fill.background()
        vt = slide.shapes.add_textbox(Inches(left), Inches(2.05), Inches(2.9), Inches(0.55))
        vp = vt.text_frame.paragraphs[0]
        vp.text = val
        vp.font.size = Pt(28)
        vp.font.bold = True
        vp.font.color.rgb = TEAL_LIGHT
        vp.alignment = PP_ALIGN.CENTER
        lt = slide.shapes.add_textbox(Inches(left), Inches(2.65), Inches(2.9), Inches(0.45))
        lp = lt.text_frame.paragraphs[0]
        lp.text = lbl
        lp.font.size = Pt(11)
        lp.font.color.rgb = WHITE
        lp.alignment = PP_ALIGN.CENTER

    bullets = slide.shapes.add_textbox(Inches(0.55), Inches(3.55), Inches(12.2), Inches(3.2))
    bf = bullets.text_frame
    for i, line in enumerate([
        "xUnit tests across Auth, Task, Project, Gateway, and shared libraries.",
        "Pull requests run backend + frontend CI before merge.",
        "Docker Compose for local SQL/Mongo/Redis; scripts/start-local.ps1 for one-command startup.",
        "Docs: README, LOCAL_DEVELOPMENT.md, gateway at :5204 for all API calls.",
    ]):
        p = bf.paragraphs[0] if i == 0 else bf.add_paragraph()
        p.text = f"▸  {line}"
        p.font.size = Pt(14)
        p.font.color.rgb = SLATE
        p.space_after = Pt(6)


def slide_thanks(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, NAVY)
    add_accent_bar(slide, 0, 0.14, TEAL)

    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.75), Inches(8), Inches(1.0))
    p = tb.text_frame.paragraphs[0]
    p.text = "Thank You"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    q = slide.shapes.add_textbox(Inches(0.6), Inches(1.65), Inches(8), Inches(0.6))
    qp = q.text_frame.paragraphs[0]
    qp.text = "We welcome your questions and feedback."
    qp.font.size = Pt(18)
    qp.font.color.rgb = TEAL_LIGHT

    lbl = slide.shapes.add_textbox(Inches(0.6), Inches(2.35), Inches(5), Inches(0.35))
    lp = lbl.text_frame.paragraphs[0]
    lp.text = "PROJECT TEAM"
    lp.font.size = Pt(11)
    lp.font.bold = True
    lp.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    add_team_photos_row(slide, start_left=0.55, photo_top=2.75, name_color=WHITE)

    contact = slide.shapes.add_textbox(Inches(0.6), Inches(6.55), Inches(12), Inches(0.4))
    cp = contact.text_frame.paragraphs[0]
    cp.text = "Task Management System  •  .NET 10 Microservices  •  React"
    cp.font.size = Pt(11)
    cp.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_title_slide(prs)
    slide_overview(prs)
    slide_problem_solution(prs)
    slide_architecture(prs)
    slide_tech_stack(prs)
    slide_services(prs)
    slide_request_flow(prs)
    slide_rbac(prs)
    slide_demo_flow(prs)
    slide_quality(prs)
    slide_thanks(prs)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
