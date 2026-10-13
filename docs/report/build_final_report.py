"""
Generate 4–5 page final project report (Word).
Requires: pip install python-docx matplotlib

Run from repo root:
  python docs/report/build_final_report.py
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "TaskManagementSystem_Final_Report.docx"

TEAM = [
    ("Denish Murawala", "Team Lead — Backend & architecture"),
    ("Avinash Kanaujia", "Frontend & user experience"),
    ("Anushtup Ganguly", "Microservices & integration"),
]

TEST_COUNTS = [
    ("Project", 51),
    ("Task", 38),
    ("Auth", 6),
    ("Dashboard", 9),
    ("File", 7),
    ("Comment", 4),
    ("Notification", 4),
    ("AI", 3),
    ("Shared", 5),
    ("Gateway", 1),
]


def ensure_assets() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    _draw_architecture_diagram(ASSETS / "architecture.png")
    _draw_test_chart(ASSETS / "unit_tests_by_service.png")
    _draw_stack_chart(ASSETS / "technology_mix.png")


def _draw_architecture_diagram(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(10, 5.5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis("off")
    fig.patch.set_facecolor("white")

    def box(x, y, w, h, text, fc="#E8F4F8", ec="#1E3A5F"):
        rect = mpatches.FancyBboxPatch(
            (x, y), w, h, boxstyle="round,pad=0.02", linewidth=1.2, edgecolor=ec, facecolor=fc
        )
        ax.add_patch(rect)
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=9, color="#0F2A44")

    box(4.0, 4.85, 2.0, 0.75, "React SPA\n(Vite, :5173)", fc="#D1FAE5", ec="#0D9488")
    box(3.55, 3.35, 2.9, 0.7, "API Gateway (Ocelot)\n:5204 — JWT validation", fc="#DBEAFE", ec="#2563EB")
    services = [
        "Auth", "Project", "Task", "Comment",
        "File", "Dashboard", "Notification", "AI",
    ]
    xs = [0.35 + i * 1.18 for i in range(8)]
    for i, (x, name) in enumerate(zip(xs, services)):
        box(x, 1.85, 1.05, 0.65, name, fc="#F8FAFC", ec="#64748B")
        ax.annotate(
            "", xy=(x + 0.52, 1.85), xytext=(4.95, 3.35),
            arrowprops=dict(arrowstyle="-", color="#94A3B8", lw=0.8, connectionstyle="arc3,rad=0.1"),
        )
    box(0.5, 0.35, 2.2, 0.7, "SQL Server\n(Users, Projects, Tasks)", fc="#FEF3C7", ec="#D97706")
    box(3.5, 0.35, 2.0, 0.7, "MongoDB\n(Comments)", fc="#FEF3C7", ec="#D97706")
    box(6.2, 0.35, 2.0, 0.7, "Local disk / Blob\n(Attachments)", fc="#FEF3C7", ec="#D97706")
    ax.annotate(
        "", xy=(1.6, 1.05), xytext=(1.0, 1.85),
        arrowprops=dict(arrowstyle="->", color="#64748B", lw=1),
    )
    ax.annotate(
        "", xy=(4.5, 1.05), xytext=(4.0, 1.85),
        arrowprops=dict(arrowstyle="->", color="#64748B", lw=1),
    )
    ax.annotate(
        "", xy=(7.2, 1.05), xytext=(6.5, 1.85),
        arrowprops=dict(arrowstyle="->", color="#64748B", lw=1),
    )
    ax.annotate(
        "", xy=(5.0, 4.85), xytext=(5.0, 4.05),
        arrowprops=dict(arrowstyle="->", color="#0D9488", lw=1.5),
    )
    ax.set_title("High-level system architecture", fontsize=12, fontweight="bold", color="#0F2A44", pad=12)
    fig.tight_layout()
    fig.savefig(path, dpi=160, bbox_inches="tight")
    plt.close(fig)


def _draw_test_chart(path: Path) -> None:
    labels = [t[0] for t in TEST_COUNTS]
    values = [t[1] for t in TEST_COUNTS]
    colors = ["#0D9488" if v >= 20 else "#3B82F6" for v in values]
    fig, ax = plt.subplots(figsize=(7, 3.8))
    bars = ax.barh(labels, values, color=colors, edgecolor="#E2E8F0")
    ax.set_xlabel("Number of unit tests (xUnit)", fontsize=10)
    ax.set_title("Automated test distribution by service", fontsize=11, fontweight="bold", color="#0F2A44")
    ax.invert_yaxis()
    for bar, v in zip(bars, values):
        ax.text(bar.get_width() + 0.5, bar.get_y() + bar.get_height() / 2, str(v), va="center", fontsize=9)
    ax.set_xlim(0, max(values) + 8)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.tight_layout()
    fig.savefig(path, dpi=160, bbox_inches="tight")
    plt.close(fig)


def _draw_stack_chart(path: Path) -> None:
    fig, ax = plt.subplots(figsize=(5, 4))
    sizes = [38, 45, 17]
    labels = ["Backend (.NET 10)", "Frontend (React)", "Data & DevOps"]
    colors = ["#1E3A5F", "#0D9488", "#F59E0B"]
    wedges, texts, autotexts = ax.pie(
        sizes, labels=labels, autopct="%1.0f%%", startangle=90, colors=colors,
        textprops={"fontsize": 9},
    )
    for t in autotexts:
        t.set_color("white")
        t.set_fontweight("bold")
    ax.set_title("Approximate effort split across layers", fontsize=11, fontweight="bold", color="#0F2A44")
    fig.tight_layout()
    fig.savefig(path, dpi=160, bbox_inches="tight")
    plt.close(fig)


def _set_document_defaults(doc: Document) -> None:
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Times New Roman"
    font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = RGBColor(0x0F, 0x2A, 0x44)


def _body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    doc.add_paragraph()


def _caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
        r.italic = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)


def _cover_page(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("TASK MANAGEMENT SYSTEM")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = RGBColor(0x0F, 0x2A, 0x44)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("A Microservices-Based Collaboration Platform\n(.NET 10, React, Docker)")
    sr.font.size = Pt(14)
    sr.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph()
    doc.add_paragraph()

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run("Final Project Report")
    mr.bold = True
    mr.font.size = Pt(16)

    doc.add_paragraph()
    team_p = doc.add_paragraph()
    team_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lines = ["Submitted by:", ""]
    for name, role in TEAM:
        lines.append(name)
        lines.append(role)
        lines.append("")
    tr = team_p.add_run("\n".join(lines))
    tr.font.size = Pt(12)

    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run(f"\n{date.today().strftime('%B %Y')}")
    fr.font.size = Pt(12)
    fr.italic = True

    doc.add_page_break()


def build_report() -> None:
    ensure_assets()
    doc = Document()
    _set_document_defaults(doc)

    # Margins ~ 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    _cover_page(doc)

    # --- Page 1 content: Abstract + Introduction ---
    _heading(doc, "Abstract", 1)
    _body(
        doc,
        "Modern software teams rarely struggle because they lack tools—they struggle because work is scattered "
        "across chats, spreadsheets, and half-finished tickets. Our Task Management System addresses that gap with "
        "a single web application where projects, sprints, Kanban boards, comments, and file attachments live together. "
        "The solution is deliberately built as a set of small .NET microservices behind an Ocelot API gateway, with "
        "a React single-page application as the only public client. JWT-based authentication separates administrators "
        "from regular users, while Docker Compose makes it practical to run SQL Server, MongoDB, and all services on "
        "a developer laptop without cloud subscriptions. This report summarises the motivation, design choices, "
        "implementation highlights, and how we validated the system before demonstration.",
    )

    _heading(doc, "1. Introduction", 1)
    _body(
        doc,
        "The project started from a simple classroom brief: demonstrate real-world .NET skills—not only REST APIs, "
        "but also service boundaries, shared security, and a usable front end. We chose task management because "
        "every reviewer immediately understands the domain: create a project, plan a sprint, assign work, and track "
        "progress on a board. That clarity let us spend more time on architecture than on explaining exotic business rules.",
    )
    _body(
        doc,
        "Over the development cycle we organised work around three roles in the team. Denish focused on gateway "
        "routing, auth, and cross-service rules (for example ensuring assignees can open tasks even when they are "
        "not yet formal project members). Avinash shaped the React UI—Kanban drag-and-drop, forms, and the overall "
        "workflow for admins versus everyday users. Anushtup connected services end to end: Docker scripts, health "
        "checks, notification hubs, and keeping local ports consistent with Ocelot configuration.",
    )

    _heading(doc, "1.1 Objectives", 2)
    objectives = [
        "Deliver a working SPA that talks to one gateway URL in development and production-style setups.",
        "Split ownership by bounded context: identity, projects/teams/sprints, tasks/subtasks, comments, files, dashboard metrics, notifications, and optional AI chat.",
        "Enforce role-based access so admins manage structure while users execute day-to-day work safely.",
        "Provide repeatable local startup (scripts + Compose) and automated tests so regressions are caught early.",
    ]
    for item in objectives:
        doc.add_paragraph(item, style="List Bullet")

    # --- Page 2: Requirements + Architecture ---
    _heading(doc, "2. Requirements and System Design", 1)
    _body(
        doc,
        "Functional requirements were captured in plain language first, then mapped to services. Administrators "
        "create teams and projects, define sprints, invite members, and assign tasks. Both admins and authorised "
        "users create tasks, move cards across statuses (Todo → In Progress → Testing → Done), complete subtasks, "
        "comment on work, and upload attachments. The dashboard aggregates counts for quick situational awareness, "
        "and SignalR pushes notification events to the browser.",
    )

    _heading(doc, "2.1 Role matrix", 2)
    _add_table(
        doc,
        ["Capability", "Admin", "User (with project access)"],
        [
            ["Create / delete teams & projects", "Yes", "No"],
            ["Manage sprints & team membership", "Yes", "No"],
            ["Assign tasks to members", "Yes", "No"],
            ["Kanban board & task editing", "Yes", "Yes"],
            ["Comments & file uploads", "Yes", "Yes"],
            ["Dashboard & notifications", "Yes", "Yes"],
        ],
    )

    _heading(doc, "2.2 Architecture", 2)
    _body(
        doc,
        "Clients never call microservices directly. The React app stores a JWT after login and sends all API traffic "
        "to the gateway on port 5204. Ocelot forwards paths such as /api/auth, /api/projects, /api/tasks, and "
        "/api/chat to the correct downstream host. Auth issues tokens; Project owns teams, projects, and sprints; "
        "Task owns tasks and subtasks; Comment persists threads in MongoDB; File stores binaries locally in "
        "development; Notification exposes REST plus hubs; Dashboard reads aggregated data; AI proxies prompts to "
        "Gemini when an API key is configured.",
    )
    doc.add_picture(str(ASSETS / "architecture.png"), width=Inches(6.2))
    _caption(doc, "Figure 1 — Logical architecture (development topology)")

    # --- Page 3: Technology + Implementation ---
    _heading(doc, "3. Technology Stack and Implementation", 1)
    _add_table(
        doc,
        ["Layer", "Technologies", "Notes"],
        [
            ["Frontend", "React 19, Vite, Tailwind CSS, Axios", "Kanban via @hello-pangea/dnd; forms with React Hook Form + Zod"],
            ["API edge", "ASP.NET Core + Ocelot", "Single public HTTP entry; CORS aligned with Vite dev server"],
            ["Services", ".NET 10 Web APIs", "EF Core on SQL where relational data fits; Mongo driver for comments"],
            ["Real-time", "SignalR", "Notification bell updates without full page refresh"],
            ["Local infra", "Docker Compose", "SQL Server + MongoDB; File service uses disk storage in Development"],
            ["Quality", "xUnit, Vitest, Playwright, GitHub Actions", "CI builds backend and frontend on pull requests"],
        ],
    )

    doc.add_picture(str(ASSETS / "technology_mix.png"), width=Inches(3.2))
    _caption(doc, "Figure 2 — Rough effort distribution (team estimate)")

    _body(
        doc,
        "Entity Framework migrations run automatically in Development for Auth, Project, Task, and File, which kept "
        "onboarding friction low for new machines. Shared libraries centralise JWT validation helpers and current-user "
        "context so each service does not reimplement security primitives. One practical lesson we documented in the "
        "repository is assignee access: users who are only assigned to a task still need read access to that project’s "
        "board. We solved it by merging task-based project IDs into the project access check and auto-adding assignees "
        "to the project team when an admin assigns work.",
    )

    _heading(doc, "3.1 Demonstration dataset", 2)
    _body(
        doc,
        "For presentations we added an idempotent PowerShell seed script. It creates the Acme Digital team, three sample "
        "projects (including Customer Mobile App), sprints, seven tasks with subtasks, assignments, and comments. "
        "Reviewers can log in as admin@local.test or user@local.test using passwords from a local .env file—secrets "
        "never committed to source control.",
    )

    # --- Page 4: Testing + Results + Conclusion ---
    _heading(doc, "4. Testing and Results", 1)
    _body(
        doc,
        "We treated automated tests as part of the deliverable, not an afterthought. Backend controllers and domain "
        "rules are covered with xUnit and in-memory or SQLite fixtures where appropriate. The gateway has configuration "
        "tests so a renamed route is caught before the UI silently fails. On the frontend, Vitest covers components and "
        "hooks; Playwright scenarios exist for end-to-end checks when both tiers are running.",
    )

    total_tests = sum(v for _, v in TEST_COUNTS)
    doc.add_picture(str(ASSETS / "unit_tests_by_service.png"), width=Inches(5.8))
    _caption(doc, f"Figure 3 — {total_tests} backend unit tests across ten test projects")

    _body(
        doc,
        "Manual testing followed a short script: start Docker and services via start-local.ps1, seed demo data, log in "
        "as admin to create structure, then log in as a regular user to drag Kanban cards and attach a file. Health "
        "scripts probe each downstream port so we do not debug the UI when the Task service simply is not running. "
        "In class demonstration the flow—login, open Customer Mobile App, show sprint backlog, move a card to Done—"
        "completed in under three minutes.",
    )

    _heading(doc, "5. Conclusion and Future Work", 1)
    _body(
        doc,
        "We met the original goal: a credible microservices sample that still feels like one product. The gateway "
        "pattern simplified the React client, and splitting comments into MongoDB avoided forcing every discussion "
        "thread into rigid SQL schemas. If we continued the project, the roadmap already notes sensible next steps—"
        "structured email notifications, OpenTelemetry traces across services, and container deployment to Azure "
        "Container Apps or Kubernetes.",
    )
    _body(
        doc,
        "Personally, the most valuable outcome was learning to negotiate service boundaries as a team. Arguments about "
        "whether sprint logic belonged with tasks or projects were resolved by reading how each API was actually consumed "
        "in the UI. That discipline is transferable to industry work far beyond this classroom scope.",
    )

    _heading(doc, "References", 1)
    refs = [
        "Microsoft Learn — ASP.NET Core microservices architecture guidance.",
        "Ocelot documentation — API gateway configuration for .NET.",
        "Project repository — README.md, docs/LOCAL_DEVELOPMENT.md, docs/REPOSITORY_LAYOUT.md.",
        "React documentation — Vite tooling and modern SPA patterns.",
    ]
    for i, ref in enumerate(refs, 1):
        doc.add_paragraph(f"[{i}] {ref}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build_report()
